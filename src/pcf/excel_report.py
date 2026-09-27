"""Write a live Excel workbook: inputs in blue, every result as a formula so it recalculates."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = "Arial"
BLUE = Font(name=F, color="0000FF")
BOLD = Font(name=F, bold=True)
BASE = Font(name=F)
TITLE = Font(name=F, bold=True, size=13)
HEAD_FILL = PatternFill("solid", start_color="E7E6E6")
KEY_FILL = PatternFill("solid", start_color="FFF2CC")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
NUM = "#,##0.00"
PCT = "0.0%"


def _header(ws, row, labels):
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, THIN
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _font_all(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != F:
                    cell.font = Font(name=F, bold=cell.font.bold, color=cell.font.color)


def build_workbook(inv_bom, variants, site, path):
    wb = Workbook()
    vs = list(variants.index)

    # Read me
    ws = wb.active
    ws.title = "Read me"
    lines = [
        ("Battery pack product carbon footprint (cradle to gate)", TITLE),
        (f"{site['company']}, reporting period {site['reporting_period']}", BASE),
        ("", BASE),
        ("Illustrative portfolio model with literature-style values. Not real company data.", BOLD),
        ("", BASE),
        ("How to use", BOLD),
        ("Blue text = inputs you can edit (Inputs and PCF sheets). Black text = formulas.", BASE),
        ("Change any blue value and every result updates automatically.", BASE),
        ("", BASE),
        ("Method", BOLD),
        ("Declared unit: one battery pack of each variant (S, M, L).", BASE),
        ("Boundary: cradle to gate (materials, inbound transport, pack assembly). Use and end of life excluded.", BASE),
        ("Impact: climate change, GWP100, kg CO2e, in line with ISO 14067.", BASE),
        ("Site energy allocated to variants by total assembly time (volume x minutes per unit).", BASE),
        ("", BASE),
        ("Sheets", BOLD),
        ("Inputs: site energy, emission factors and variant data", BASE),
        ("Allocation: splits site electricity and gas across variants", BASE),
        ("PCF: component inventory and footprint per pack", BASE),
        ("Summary: hotspots by category", BASE),
        ("Corporate GHG: annual Scope 1, 2 and 3 roll-up and intensity", BASE),
        ("Frameworks: where each number feeds GHG Protocol, CSRD, TCFD and EU Taxonomy reporting", BASE),
    ]
    for i, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font
    ws.column_dimensions["A"].width = 100

    # Inputs
    ws = wb.create_sheet("Inputs")
    ws["A1"], ws["A1"].font = "Inputs", TITLE
    _header(ws, 3, ["Site data", "Value", "Unit"])
    site_rows = [
        ("Reporting period", site["reporting_period"], ""),
        ("Site electricity", site["site_electricity_kwh"], "kWh per year"),
        ("Site natural gas", site["site_natural_gas_kwh"], "kWh per year"),
        ("Grid emission factor (location-based)", site["grid_factor_kgco2e_per_kwh"], "kg CO2e per kWh"),
        ("Natural gas emission factor", site["gas_factor_kgco2e_per_kwh"], "kg CO2e per kWh"),
    ]
    for r, (label, val, unit) in enumerate(site_rows, 4):
        ws.cell(row=r, column=1, value=label).font = BASE
        c = ws.cell(row=r, column=2, value=val)
        c.font = BLUE
        if isinstance(val, float):
            c.number_format = "0.000"
        elif isinstance(val, int):
            c.number_format = "#,##0"
        ws.cell(row=r, column=3, value=unit).font = BASE
    # Inputs!B5 electricity, B6 gas, B7 grid factor, B8 gas factor
    _header(ws, 10, ["Variant", "Capacity (kWh)", "Annual volume (units)", "Assembly time (min per unit)"])
    for r, v in enumerate(vs, 11):
        ws.cell(row=r, column=1, value=v).font = BASE
        for c, col in enumerate(["capacity_kwh", "annual_volume", "assembly_minutes"], 2):
            cell = ws.cell(row=r, column=c, value=variants.loc[v, col].item())
            cell.font = BLUE
            cell.number_format = "#,##0.0" if col == "capacity_kwh" else "#,##0"
    _widths(ws, [40, 18, 22, 26])
    vrow = {v: 11 + i for i, v in enumerate(vs)}  # Inputs rows for each variant

    # Allocation
    ws = wb.create_sheet("Allocation")
    ws["A1"], ws["A1"].font = "Site energy allocation by assembly time", TITLE
    _header(ws, 3, ["Variant", "Annual volume", "Assembly min per unit", "Total assembly time (min)",
                    "Time share", "Electricity (kWh per unit)", "Natural gas (kWh per unit)"])
    first, last = 4, 3 + len(vs)
    for r, v in enumerate(vs, 4):
        ir = vrow[v]
        ws.cell(row=r, column=1, value=f"=Inputs!A{ir}")
        ws.cell(row=r, column=2, value=f"=Inputs!C{ir}").number_format = "#,##0"
        ws.cell(row=r, column=3, value=f"=Inputs!D{ir}").number_format = "#,##0"
        ws.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format = "#,##0"
        ws.cell(row=r, column=5, value=f"=D{r}/SUM($D${first}:$D${last})").number_format = PCT
        ws.cell(row=r, column=6, value=f"=IF(B{r}=0,0,Inputs!$B$5*E{r}/B{r})").number_format = NUM
        ws.cell(row=r, column=7, value=f"=IF(B{r}=0,0,Inputs!$B$6*E{r}/B{r})").number_format = NUM
    tr = last + 1
    ws.cell(row=tr, column=1, value="Total").font = BOLD
    ws.cell(row=tr, column=4, value=f"=SUM(D{first}:D{last})").number_format = "#,##0"
    ws.cell(row=tr, column=5, value=f"=SUM(E{first}:E{last})").number_format = PCT
    _widths(ws, [10, 15, 20, 24, 12, 24, 24])
    arow = {v: 4 + i for i, v in enumerate(vs)}

    # PCF
    ws = wb.create_sheet("PCF")
    ws["A1"], ws["A1"].font = "Product carbon footprint per pack (kg CO2e)", TITLE
    amt_cols = {v: 8 + i for i, v in enumerate(vs)}
    gwp_cols = {v: 8 + len(vs) + i for i, v in enumerate(vs)}
    _header(ws, 3, ["Component", "Category", "GHG scope", "Data type", "Unit", "Emission factor (kg CO2e per unit)",
                    "Factor source"] + [f"Amount {v}" for v in vs] + [f"kg CO2e {v}" for v in vs])
    r = 4
    for _, row in inv_bom.iterrows():
        for c, key in enumerate(["component", "category", "ghg_scope", "data_type", "unit"], 1):
            ws.cell(row=r, column=c, value=row[key]).font = BASE
        f = ws.cell(row=r, column=6, value=float(row["emission_factor"]))
        f.font, f.number_format = BLUE, "0.00"
        ws.cell(row=r, column=7, value=row["factor_source"]).font = BASE
        for v in vs:
            a = ws.cell(row=r, column=amt_cols[v], value=float(row[v]))
            a.font, a.number_format = BLUE, "#,##0.00"
        r += 1
    energy_rows = [
        ("Assembly electricity (allocated)", "Scope 2", "=Inputs!$B$7", "F", "Location-based grid factor"),
        ("Assembly heating, natural gas (allocated)", "Scope 1", "=Inputs!$B$8", "G", "Natural gas combustion factor"),
    ]
    for name, scope, factor, acol, src in energy_rows:
        vals = [name, "Assembly energy", scope, "Primary", "kWh"]
        for c, val in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=val).font = BASE
        ws.cell(row=r, column=6, value=factor).number_format = "0.00"
        ws.cell(row=r, column=7, value=src).font = BASE
        for v in vs:
            ws.cell(row=r, column=amt_cols[v], value=f"=Allocation!{acol}{arow[v]}").number_format = "#,##0.00"
        r += 1
    last_inv = r - 1
    for rr in range(4, last_inv + 1):
        for v in vs:
            ac = get_column_letter(amt_cols[v])
            ws.cell(row=rr, column=gwp_cols[v], value=f"=$F{rr}*{ac}{rr}").number_format = NUM
    tot, cap, per_kwh, share = last_inv + 2, last_inv + 3, last_inv + 4, last_inv + 5
    labels = {tot: "Total PCF (kg CO2e per pack)", cap: "Pack capacity (kWh)",
              per_kwh: "PCF per kWh capacity (kg CO2e per kWh)", share: "Battery cell share of PCF"}
    for rr, lab in labels.items():
        ws.cell(row=rr, column=1, value=lab).font = BOLD
    for v in vs:
        col = get_column_letter(gwp_cols[v])
        ws[f"{col}{tot}"] = f"=SUM({col}4:{col}{last_inv})"
        ws[f"{col}{cap}"] = f"=Inputs!B{vrow[v]}"
        ws[f"{col}{per_kwh}"] = f"=IF({col}{cap}=0,0,{col}{tot}/{col}{cap})"
        ws[f"{col}{share}"] = f'=IF({col}{tot}=0,0,SUMIF($B$4:$B${last_inv},"Battery cells",{col}4:{col}{last_inv})/{col}{tot})'
        for rr, fmt in [(tot, NUM), (cap, "0.0"), (per_kwh, NUM), (share, PCT)]:
            ws[f"{col}{rr}"].number_format = fmt
            ws[f"{col}{rr}"].font = BOLD
            if rr in (tot, share):
                ws[f"{col}{rr}"].fill = KEY_FILL
    _widths(ws, [40, 16, 15, 11, 7, 18, 45] + [11] * len(vs) + [13] * len(vs))
    ws.freeze_panes = "B4"

    # Summary (hotspots by category)
    ws = wb.create_sheet("Summary")
    ws["A1"], ws["A1"].font = "Hotspots by category (kg CO2e per pack)", TITLE
    cats = list(dict.fromkeys(inv_bom["category"])) + ["Assembly energy"]
    _header(ws, 3, ["Category"] + [f"Pack {v}" for v in vs] + [f"Share {v}" for v in vs])
    for i, cat in enumerate(cats):
        rr = 4 + i
        ws.cell(row=rr, column=1, value=cat).font = BASE
        for j, v in enumerate(vs):
            g = get_column_letter(gwp_cols[v])
            ws.cell(row=rr, column=2 + j, value=f"=SUMIF(PCF!$B$4:$B${last_inv},$A{rr},PCF!{g}$4:{g}${last_inv})").number_format = NUM
            ws.cell(row=rr, column=2 + len(vs) + j,
                    value=f"=IF(PCF!{g}${tot}=0,0,{get_column_letter(2 + j)}{rr}/PCF!{g}${tot})").number_format = PCT
    sr = 4 + len(cats)
    ws.cell(row=sr, column=1, value="Total").font = BOLD
    for j in range(len(vs)):
        col = get_column_letter(2 + j)
        c = ws.cell(row=sr, column=2 + j, value=f"=SUM({col}4:{col}{sr - 1})")
        c.number_format, c.font = NUM, BOLD
    _widths(ws, [22] + [12] * (2 * len(vs)))

    # Corporate GHG
    ws = wb.create_sheet("Corporate GHG")
    ws["A1"], ws["A1"].font = f"Annual GHG inventory roll-up, {site['reporting_period']} (illustrative)", TITLE
    _header(ws, 3, ["GHG scope"] + [f"Per unit {v} (kg CO2e)" for v in vs] + ["Annual total (t CO2e)", "Share"])
    scopes = ["Scope 1", "Scope 2", "Scope 3 Cat 1", "Scope 3 Cat 4"]
    tcol = get_column_letter(2 + len(vs))
    for i, sc in enumerate(scopes):
        rr = 4 + i
        ws.cell(row=rr, column=1, value=sc).font = BASE
        parts = []
        for j, v in enumerate(vs):
            g = get_column_letter(gwp_cols[v])
            ws.cell(row=rr, column=2 + j,
                    value=f"=SUMIF(PCF!$C$4:$C${last_inv},$A{rr},PCF!{g}$4:{g}${last_inv})").number_format = NUM
            parts.append(f"{get_column_letter(2 + j)}{rr}*Inputs!$C${vrow[v]}")
        ws.cell(row=rr, column=2 + len(vs), value="=(" + "+".join(parts) + ")/1000").number_format = NUM
    tr = 4 + len(scopes)
    ws.cell(row=tr, column=1, value="Total").font = BOLD
    c = ws.cell(row=tr, column=2 + len(vs), value=f"=SUM({tcol}4:{tcol}{tr - 1})")
    c.number_format, c.font, c.fill = NUM, BOLD, KEY_FILL
    for rr in range(4, tr):
        ws.cell(row=rr, column=3 + len(vs), value=f"=IF(${tcol}${tr}=0,0,{tcol}{rr}/${tcol}${tr})").number_format = PCT
    cap_expr = "+".join(f"Inputs!B{vrow[v]}*Inputs!C{vrow[v]}" for v in vs)
    extra = [
        ("Battery capacity produced (MWh)", f"=({cap_expr})/1000", "#,##0.0"),
        ("GHG intensity (kg CO2e per kWh produced)", f"=IF({tcol}{tr + 2}=0,0,{tcol}{tr}/{tcol}{tr + 2})", NUM),
        ("Check: Scope 1 from site gas (t CO2e)", "=Inputs!B6*Inputs!B8/1000", NUM),
        ("Check: Scope 2 from site electricity (t CO2e)", "=Inputs!B5*Inputs!B7/1000", NUM),
    ]
    for k, (lab, formula, fmt) in enumerate(extra):
        rr = tr + 2 + k
        ws.cell(row=rr, column=1, value=lab).font = BOLD if k < 2 else BASE
        c = ws.cell(row=rr, column=2 + len(vs), value=formula)
        c.number_format = fmt
    ws.cell(row=tr + 7, column=1, value="The two checks should match the Scope 1 and Scope 2 totals above, "
                                         "which proves the allocation is complete.").font = Font(name=F, italic=True)
    _widths(ws, [44] + [20] * len(vs) + [20, 10])

    # Frameworks
    ws = wb.create_sheet("Frameworks")
    ws["A1"], ws["A1"].font = "Where each result feeds reporting frameworks (simplified learning overview)", TITLE
    _header(ws, 3, ["Result", "GHG Protocol", "CSRD (ESRS E1)", "TCFD", "EU Taxonomy", "Transition risk angle"])
    rows = [
        ("Scope 1 (site gas)", "Corporate Standard, Scope 1", "E1-6 gross Scope 1", "Metrics and targets",
         "", "Carbon pricing on heating fuels (planned EU ETS2) raises gas costs"),
        ("Scope 2 (site electricity)", "Scope 2 Guidance: location and market-based", "E1-6 Scope 2, both methods",
         "Metrics and targets", "", "Electricity prices; renewable PPAs as a lever"),
        ("Scope 3 Cat 1 (purchased goods)", "Scope 3 Standard, Category 1", "E1-6 significant Scope 3 categories",
         "Metrics (Scope 3 where material)", "", "Supplier emissions dominate; customers request low-carbon cells"),
        ("Scope 3 Cat 4 (inbound transport)", "Scope 3 Standard, Category 4", "E1-6 Scope 3", "Metrics",
         "", "Freight decarbonisation and fuel costs"),
        ("PCF per pack", "Product Standard / ISO 14067", "Supports E1 transition plan and targets",
         "Strategy", "Activity 3.4 Manufacture of batteries",
         "EU Battery Regulation carbon footprint declarations for certain battery categories"),
        ("GHG intensity", "", "E1-6 intensity is per net revenue; per kWh is an extra product metric",
         "Metrics", "", "Benchmarking against competitors and customers' targets"),
        ("Aluminium housing", "Scope 3 Cat 1", "", "Risk management", "",
         "CBAM covers aluminium imports; recycled aluminium reduces cost and footprint exposure"),
    ]
    for i, row in enumerate(rows, 4):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=val)
            cell.font = BASE
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=4 + len(rows) + 1, column=1,
            value="Note: CSRD scope and timelines were revised by the EU Omnibus simplification. "
                  "Always check the latest regulatory text.").font = Font(name=F, italic=True)
    _widths(ws, [30, 30, 34, 24, 26, 48])

    _font_all(wb)
    wb.save(path)
